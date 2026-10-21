"""UAT outputs: the Excel sheet (summary, test cases, functional matrix, accounts) and the evidence page with the videos."""

import html
import os
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="0B6E4F")
PASS = PatternFill("solid", fgColor="E5F4E8")
FAIL = PatternFill("solid", fgColor="FDE8E7")
THIN = Side(style="thin", color="D0D7DE")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def table(ws, headers, rows, widths):
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font, cell.fill, cell.alignment, cell.border = Font(bold=True, color="FFFFFF"), HEAD, WRAP, BOX
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.alignment, cell.border = WRAP, BOX
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=ws.max_row - len(rows) + 1 if rows else 2, column=1)


def color_status(ws, column):
    for row in ws.iter_rows(min_row=2):
        cell = row[column - 1]
        if cell.value == "Đạt":
            cell.fill = PASS
        elif cell.value == "Không đạt":
            cell.fill = FAIL


def personas(results):
    groups = OrderedDict()
    for r in results:
        groups.setdefault((r["role"], r["account"]), []).append(r)
    return groups


def write_xlsx(out, run_id, results, matrix, videos, accounts, sections, labels, expected):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng hợp"
    passed = sum(1 for r in results if r["status"] == "Đạt")
    ws.append([f"UAT TLC HRM trên Frappe HR - lần chạy {run_id}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Kết quả: {passed}/{len(results)} trường hợp Đạt. Môi trường: k3s ai1, site tlc-hr.local, image tlc-hr:{os.environ.get('IMAGE_TAG', '?')}, dữ liệu demo TLC (60 nhân sự hư cấu)."])
    ws.append([])
    rows = []
    for (role, account), items in personas(results).items():
        ok = sum(1 for r in items if r["status"] == "Đạt")
        vids = sorted({r["video"] for r in items if r["video"]})
        rows.append([role, account, len(items), ok, len(items) - ok, "\n".join(vids) or "(bước API, không ghi hình)"])
    table(ws, ["Vai trò", "Tài khoản", "Số TC", "Đạt", "Không đạt", "Video bằng chứng"], rows, [44, 26, 8, 8, 10, 46])

    ws = wb.create_sheet("Kịch bản UAT")
    table(ws, ["Mã TC", "Vai trò", "Tài khoản", "Chức năng", "Các bước thực hiện", "Dữ liệu test", "Kết quả mong đợi", "Kết quả thực tế",
               "Trạng thái", "Video", "Mốc thời gian (mm:ss)", "Ảnh chụp", "Thời điểm chạy"],
          [[r["id"], r["role"], r["account"], r["function"], r["steps"], r["data"], r["expected"], r["actual"], r["status"],
            r["video"] or "", r["time"], r["shot"], r["run_at"]] for r in results],
          [9, 30, 22, 30, 40, 26, 40, 46, 11, 30, 12, 22, 19])
    color_status(ws, 9)

    ws = wb.create_sheet("Ma trận chức năng")
    headers = ["Tài khoản", "Vai trò"] + [labels[s] for s in sections] + ["Khớp mong đợi"]
    rows = []
    for account, info in accounts.items():
        shown = set(matrix.get(account, []))
        cells = []
        for s in sections:
            want, got = s in expected.get(account, set()), s in shown
            cells.append(("✔ hiện" if got else "— ẩn") + ("" if want == got else " (SAI)"))
        rows.append([f"{account}@demo.example", info["role"]] + cells + ["Đạt" if shown == expected.get(account, set()) else "Không đạt"])
    table(ws, headers, rows, [22, 40] + [16] * len(sections) + [14])
    color_status(ws, len(headers))

    ws = wb.create_sheet("Tài khoản test")
    table(ws, ["Tài khoản", "Họ tên", "Vai trò", "Khối", "Người duyệt nghỉ phép", "Mật khẩu"],
          [[f"{a}@demo.example", i["name"] or "", i["role"], i["workforce"], i["approver"] or "(không có - CEO)",
            "Mật khẩu demo chung: Secret tlc-hr, khóa demo-password"] for a, i in accounts.items()],
          [24, 22, 44, 10, 24, 46])

    name = f"UAT_TLC_HRM_{run_id}.xlsx"
    wb.save(out / name)
    return name


def write_html(out, run_id, results, videos, xlsx):
    passed = sum(1 for r in results if r["status"] == "Đạt")
    e = html.escape
    cards = []
    for v in videos:
        rows = "".join(
            f"<tr class='{'ok' if r['status'] == 'Đạt' else 'bad'}'><td>{e(r['id'])}</td><td>{e(r['function'])}</td><td>{e(r['time'])}</td>"
            f"<td>{e(r['expected'])}</td><td>{e(r['actual'])}</td><td><b>{e(r['status'])}</b></td>"
            f"<td>{'<a href=' + chr(39) + e(r['shot']) + chr(39) + '>ảnh</a>' if r['shot'] else ''}</td></tr>"
            for r in results if r["video"] == v["file"])
        cards.append(f"""<section><h2>{e(v['code'])} · {e(v['role'])} <small>{e(v['account'])}@demo.example · {e(v['device'])}</small></h2>
<video controls preload="metadata" src="{e(v['file'])}"></video>
<table><tr><th>Mã</th><th>Chức năng</th><th>Mốc</th><th>Mong đợi</th><th>Thực tế</th><th>KQ</th><th></th></tr>{rows}</table></section>""")
    support = "".join(f"<tr class='{'ok' if r['status'] == 'Đạt' else 'bad'}'><td>{e(r['id'])}</td><td>{e(r['account'])}</td>"
                      f"<td>{e(r['function'])}</td><td>{e(r['actual'])}</td><td><b>{e(r['status'])}</b></td></tr>"
                      for r in results if not r["video"])
    page = f"""<!doctype html><html lang="vi"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>UAT TLC HRM {e(run_id)}</title>
<style>body{{font:15px/1.45 system-ui,sans-serif;max-width:1200px;margin:24px auto;padding:0 16px;color:#1f2933}}
h1{{margin:0 0 4px}} .sum{{font-size:18px;margin:8px 0 20px}} section{{border:1px solid #e3e7eb;border-radius:12px;padding:16px;margin:0 0 20px}}
h2{{font-size:18px;margin:0 0 10px}} h2 small{{color:#66727f;font-weight:400;font-size:14px}} video{{width:100%;max-height:620px;background:#000;border-radius:8px}}
table{{width:100%;border-collapse:collapse;font-size:13px;margin-top:10px}} th,td{{border-bottom:1px solid #e3e7eb;padding:6px;text-align:left;vertical-align:top}}
tr.ok td:nth-child(6),tr.ok td:nth-child(5):last-child{{color:#1e7b34}} tr.bad{{background:#fde8e7}} a{{color:#0b6e4f}}</style>
<h1>UAT TLC HRM trên Frappe HR</h1>
<div class="sum">Lần chạy {e(run_id)} · <b>{passed}/{len(results)}</b> trường hợp Đạt · image tlc-hr:{e(os.environ.get('IMAGE_TAG', '?'))} · <a href="{e(xlsx)}">Tải sheet UAT (Excel)</a> · <a href="results.json">results.json</a></div>
<section><h2>Bước hỗ trợ (API, không ghi hình)</h2><table><tr><th>Mã</th><th>Tài khoản</th><th>Chức năng</th><th>Thực tế</th><th>KQ</th></tr>{support}</table></section>
{''.join(cards)}
</html>"""
    (out / "index.html").write_text(page, encoding="utf-8")
